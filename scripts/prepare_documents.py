#!/usr/bin/env python3
"""Extract supported document text and emit provenance-preserving JSONL chunks."""
from __future__ import annotations

import argparse
import hashlib
import json
import sys
from pathlib import Path
from typing import Iterable

SUPPORTED = {".txt", ".md", ".markdown", ".pdf", ".docx"}


def read_segments(path: Path) -> list[tuple[str, str]]:
    """Return (location, text) segments; PDF segments preserve page numbers."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        try:
            import fitz  # PyMuPDF
        except ImportError as exc:
            raise RuntimeError("PDF extraction needs PyMuPDF: pip install PyMuPDF") from exc
        segments = []
        with fitz.open(path) as pdf:
            for index, page in enumerate(pdf, start=1):
                text = page.get_text("text").strip()
                if text:
                    segments.append((f"page {index}", text))
        return segments
    if suffix == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("DOCX extraction needs python-docx: pip install python-docx") from exc
        doc = Document(path)
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        # Include table content, which commonly carries important definitions and values.
        for table in doc.tables:
            for row in table.rows:
                cell_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if cell_text:
                    paragraphs.append(cell_text)
        return [("document", "\n\n".join(paragraphs))] if paragraphs else []
    text = path.read_text(encoding="utf-8", errors="replace").strip()
    return [("document", text)] if text else []


def chunk_words(text: str, size: int, overlap: int) -> Iterable[str]:
    words = text.split()
    if not words:
        return
    stride = size - overlap
    for start in range(0, len(words), stride):
        chunk = " ".join(words[start : start + size]).strip()
        if chunk:
            yield chunk
        if start + size >= len(words):
            break


def build_chunks(input_path: Path, chunk_size: int, overlap: int) -> list[dict]:
    root = input_path if input_path.is_dir() else input_path.parent
    files = sorted(p for p in (input_path.rglob("*") if input_path.is_dir() else [input_path])
                   if p.is_file() and p.suffix.lower() in SUPPORTED)
    records: list[dict] = []
    for path in files:
        source = path.relative_to(root).as_posix()
        for base_location, text in read_segments(path):
            for index, chunk in enumerate(chunk_words(text, chunk_size, overlap), start=1):
                location = base_location if base_location.startswith("page ") else f"{base_location}, chunk {index}"
                digest = hashlib.sha256(f"{source}\0{location}\0{chunk}".encode("utf-8")).hexdigest()[:20]
                records.append({
                    "chunk_id": f"{digest}",
                    "source": source,
                    "location": location,
                    "chunk_index": index,
                    "text": chunk,
                })
    return records


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, type=Path, help="A file or directory of source documents")
    parser.add_argument("--output", required=True, type=Path, help="Output JSONL path")
    parser.add_argument("--chunk-words", type=int, default=1000)
    parser.add_argument("--overlap-words", type=int, default=120)
    parser.add_argument("--overwrite", action="store_true", help="Replace an existing output file")
    args = parser.parse_args()
    if not args.input.exists():
        parser.error(f"input does not exist: {args.input}")
    if args.chunk_words < 1 or args.overlap_words < 0 or args.overlap_words >= args.chunk_words:
        parser.error("require chunk-words >= 1 and 0 <= overlap-words < chunk-words")
    if args.output.exists() and not args.overwrite:
        parser.error(f"output already exists: {args.output}; pass --overwrite to replace it")
    try:
        records = build_chunks(args.input, args.chunk_words, args.overlap_words)
    except (RuntimeError, OSError) as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 2
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("w", encoding="utf-8") as handle:
        for record in records:
            handle.write(json.dumps(record, ensure_ascii=False) + "\n")
    print(json.dumps({"documents": len({r['source'] for r in records}), "chunks": len(records), "output": str(args.output)}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
